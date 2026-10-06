#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel 报表格式生成引擎
=======================

本模块只负责一件事：把一份「报表定义」(ReportSpec) 渲染成一个排版好的、
带格式但内容为空的 Excel 工作簿，供使用者后续填入数据。

它不关心具体是哪种报表——财务状况表、工资表、盘点表的差异全部体现在
report_specs.py 里的数据定义中。想新增一种报表，只需要往那边加一个
ReportSpec，本文件一行都不用改。

主要对外接口：
    build_workbook(spec, ...)  -> openpyxl.Workbook   生成工作簿对象
    workbook_to_bytes(wb)      -> bytes               转成 xlsx 二进制，供下载
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties


# ---------------------------------------------------------------------------
# 样式常量（想换配色 / 字体，只改这里）
# ---------------------------------------------------------------------------

FONT_NAME = "微软雅黑"          # 中文报表通用字体，Windows 自带

COLOR_HEADER_BG = "2F5597"      # 表头底色：深蓝
COLOR_HEADER_FG = "FFFFFF"      # 表头字色：白
COLOR_SECTION_BG = "DEEAF6"     # 分类行底色：浅蓝
COLOR_TOTAL_BG = "FFF2CC"       # 合计行底色：浅黄
COLOR_STRIPE_BG = "F5F8FC"      # 明细隔行底色：极浅蓝灰
COLOR_BORDER = "9DB2CE"         # 边框颜色：灰蓝
COLOR_TITLE = "1F3864"          # 大标题字色：深蓝黑

# 各种行高（单位：磅）
HEIGHT_TITLE = 34
HEIGHT_SUBTITLE = 22
HEIGHT_HEADER = 30
HEIGHT_ITEM = 20
HEIGHT_SECTION = 22
HEIGHT_TOTAL = 22
HEIGHT_FOOTER = 26

# Excel 的「下拉列表」数据有效性，选项串最长 255 字符，超了会报错
_DV_MAX_LEN = 250


# ---------------------------------------------------------------------------
# 报表定义用的数据结构
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Column:
    """
    描述报表的一列。

    字段：
        header        : 表头文字，例如「期末余额」
        width         : 列宽（Excel 的字符宽度单位，中文一个字约等于 2）
        number_format : 数字格式。常用取值：
                          "General"     常规
                          "#,##0.00"    金额，千分位 + 两位小数
                          "#,##0"       整数，千分位
                          "0.0%"        百分比
                          "yyyy-mm-dd"  日期
                          "@"           文本（防止前导零被吃掉，如工号 0012）
        align         : 数据区水平对齐，left / center / right
        summable      : 是否参与「合计」求和。金额列设 True，
                        「行次」「备注」这类列必须设 False，否则会被错误相加
        choices       : 非空则为该列生成下拉选择框，例如 ("是", "否")
    """
    header: str
    width: float = 14.0
    number_format: str = "General"
    align: str = "center"
    summable: bool = False
    choices: tuple[str, ...] = ()


@dataclass(frozen=True)
class RowSpec:
    """
    描述「固定行报表」里的一行（例如财务状况表的每一个会计科目）。

    明细流水型报表（收支明细、盘点表等）行数不固定，不需要用到它。

    字段：
        label    : 显示在第一列的文字
        ref      : 行的唯一标识，供合计公式引用（不参与显示）
        indent   : 缩进层级，0 = 顶格，1 = 一级子项，2 = 二级子项
        kind     : "item"    普通科目行
                   "section" 分类行，加粗 + 浅蓝底
                   "total"   合计行，加粗 + 浅黄底，通常配 total_of / total_refs
        total_of : (起始行 ref, 结束行 ref)，用于「连续区间」求和，
                   引擎会自动换成真实行号，为所有 summable 列生成
                   =SUM(X起:X止) 公式。例如「流动资产合计」
        total_refs : 若干个不连续行的 ref，用于「跨段相加」，生成
                   =SUM(X行1,X行2) 公式。例如「资产总计」=
                   流动资产合计 + 非流动资产合计，这两行并不相邻
        formula  : 手写公式模板，只作用于 summable 列。模板里可以写
                   {col} 表示当前列字母、{某行ref} 表示那一行的行号。
                   例如固定资产账面价值 = 原值 - 累计折旧，写成：
                       "={col}{fa_gross}-{col}{fa_dep}"
                   优先级高于 total_of / total_refs
    """
    label: str
    ref: str = ""
    indent: int = 0
    kind: str = "item"
    total_of: tuple[str, str] | None = None
    total_refs: tuple[str, ...] = ()
    formula: str = ""


@dataclass(frozen=True)
class ReportSpec:
    """一种报表的完整定义。"""
    key: str                       # 程序内部用的唯一标识（英文）
    name: str                      # 界面上显示的中文名，例如「财务状况表」
    description: str               # 一句话说明用途
    sheet_name: str                # 工作表标签名
    default_title: str             # 表格大标题的默认文字
    columns: tuple[Column, ...]    # 列定义
    rows: tuple[RowSpec, ...] = () # 固定行定义；留空表示明细流水型
    unit_label: str = "单位：元"    # 右上角的单位说明
    has_total_row: bool = True     # 明细型报表是否在末尾追加「合计」行
    stripe: bool = True            # 明细区是否使用隔行底色
    landscape: bool = True         # 打印方向：True 横向 / False 纵向
    seq_column: int = 0            # 需要自动编号的列号（1 起）；0 表示不编号
    footer_labels: tuple[str, ...] = ("制表人：", "审核人：", "负责人：")
    remark_line: bool = True       # 是否留一行「备注：」

    @property
    def is_fixed_rows(self) -> bool:
        """True = 行结构固定（如财务状况表）；False = 明细流水型。"""
        return bool(self.rows)


# ---------------------------------------------------------------------------
# 内部工具
# ---------------------------------------------------------------------------

def _font(size: int = 11, bold: bool = False, color: str = "000000") -> Font:
    """快捷生成字体对象。"""
    return Font(name=FONT_NAME, size=size, bold=bold, color=color)


def _fill(rgb: str) -> PatternFill:
    """快捷生成纯色底纹。"""
    return PatternFill("solid", fgColor=rgb)


def _align(horizontal: str = "center", vertical: str = "center",
           indent: int = 0, wrap: bool = False) -> Alignment:
    """快捷生成对齐方式。"""
    return Alignment(horizontal=horizontal, vertical=vertical,
                     indent=indent, wrap_text=wrap)


def _prepare_sheet(spec: ReportSpec, base_bytes: bytes | None) -> tuple[Workbook, object]:
    """
    准备一个干净的工作表。

    base_bytes 为 None 时新建工作簿；否则读取用户上传的 Excel，
    删掉里面原有的所有工作表，换成一张以报表名命名的空白表。

    注意：这里操作的是内存中的副本，用户硬盘上的原文件不会被修改。
    """
    if base_bytes:
        wb = load_workbook(BytesIO(base_bytes))
        # 先用临时名建表，再删掉所有旧表，最后改名——
        # 这样即使上传的文件里已经有一张同名表也不会冲突
        ws = wb.create_sheet(title="__tmp__")
        for name in list(wb.sheetnames):
            if name != "__tmp__":
                del wb[name]
        ws.title = spec.sheet_name
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = spec.sheet_name

    # 报表是给人看的成品，隐藏网格线后更接近打印效果
    ws.sheet_view.showGridLines = False
    return wb, ws


def _write_title_block(ws, spec: ReportSpec, title: str, org: str, period: str) -> int:
    """
    写最上面的「大标题 + 副标题」两行，返回表头所在的行号。

    版面结构：
        第 1 行  居中大标题
        第 2 行  左：编制单位   中：期间   右：单位
        第 3 行  表头
    """
    n_cols = len(spec.columns)
    last_letter = get_column_letter(n_cols)

    # --- 第 1 行：大标题 ---
    ws.merge_cells(f"A1:{last_letter}1")
    cell = ws["A1"]
    cell.value = title or spec.default_title
    cell.font = _font(size=16, bold=True, color=COLOR_TITLE)
    cell.alignment = _align()
    ws.row_dimensions[1].height = HEIGHT_TITLE

    # --- 第 2 行：编制单位 / 期间 / 单位 ---
    # 按列数把这一行大致切成三段；列数太少时退化成整行居中
    left_text = f"编制单位：{org}" if org else "编制单位："
    center_text = period or ""
    right_text = spec.unit_label

    if n_cols >= 3:
        a = max(1, round(n_cols * 0.4))          # 左段结束列
        b = max(a + 1, round(n_cols * 0.72))     # 中段结束列
        segments = [
            (1, a, left_text, "left"),
            (a + 1, b, center_text, "center"),
            (b + 1, n_cols, right_text, "right"),
        ]
    else:
        merged = "    ".join(t for t in (left_text, center_text, right_text) if t)
        segments = [(1, n_cols, merged, "center")]

    for start, end, text, horizontal in segments:
        if end > start:
            ws.merge_cells(start_row=2, start_column=start,
                           end_row=2, end_column=end)
        c = ws.cell(row=2, column=start, value=text)
        c.font = _font(size=10)
        c.alignment = _align(horizontal=horizontal)
    ws.row_dimensions[2].height = HEIGHT_SUBTITLE

    return 3   # 表头在第 3 行


def _write_header(ws, spec: ReportSpec, header_row: int) -> None:
    """写表头行：深蓝底白字、加粗、居中、自动换行、加边框。"""
    for idx, col in enumerate(spec.columns, start=1):
        cell = ws.cell(row=header_row, column=idx, value=col.header)
        cell.font = _font(size=11, bold=True, color=COLOR_HEADER_FG)
        cell.fill = _fill(COLOR_HEADER_BG)
        cell.alignment = _align(wrap=True)
        cell.border = Border(left=Side("thin", color=COLOR_HEADER_FG),
                             right=Side("thin", color=COLOR_HEADER_FG),
                             top=Side("thin", color=COLOR_BORDER),
                             bottom=Side("medium", color=COLOR_BORDER))
        ws.column_dimensions[get_column_letter(idx)].width = col.width
    ws.row_dimensions[header_row].height = HEIGHT_HEADER


def _style_data_cell(cell, col: Column, *, bg: str | None = None,
                     bold: bool = False, indent: int = 0,
                     horizontal: str | None = None) -> None:
    """给数据区的单个单元格套用统一样式。"""
    cell.font = _font(size=10.5, bold=bold)
    cell.border = Border(left=Side("thin", color=COLOR_BORDER),
                         right=Side("thin", color=COLOR_BORDER),
                         top=Side("thin", color=COLOR_BORDER),
                         bottom=Side("thin", color=COLOR_BORDER))
    cell.alignment = _align(horizontal=horizontal or col.align, indent=indent)
    cell.number_format = col.number_format
    if bg:
        cell.fill = _fill(bg)


def _add_dropdowns(ws, spec: ReportSpec, first_row: int, last_row: int) -> None:
    """为带 choices 的列批量添加下拉选择框。"""
    if last_row < first_row:
        return
    for idx, col in enumerate(spec.columns, start=1):
        if not col.choices:
            continue
        # Excel 要求选项用英文逗号连接，并且整体带一对双引号
        joined = ",".join(col.choices)
        if len(joined) > _DV_MAX_LEN:
            # 选项太多会超出 Excel 限制，此时放弃下拉而不是让文件损坏
            continue
        dv = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True)
        dv.error = "请从下拉列表中选择一个选项"
        dv.errorTitle = "输入无效"
        dv.prompt = "点击右侧箭头选择"
        dv.promptTitle = col.header
        ws.add_data_validation(dv)
        letter = get_column_letter(idx)
        dv.add(f"{letter}{first_row}:{letter}{last_row}")


def _write_total_row(ws, spec: ReportSpec, row: int,
                     first_data_row: int, last_data_row: int) -> None:
    """
    在明细区下方写一行「合计」：对每个 summable 列生成 =SUM() 公式。
    """
    for idx, col in enumerate(spec.columns, start=1):
        cell = ws.cell(row=row, column=idx)
        if idx == 1:
            cell.value = "合计"
        elif col.summable and last_data_row >= first_data_row:
            letter = get_column_letter(idx)
            cell.value = f"=SUM({letter}{first_data_row}:{letter}{last_data_row})"
        _style_data_cell(cell, col, bg=COLOR_TOTAL_BG, bold=True,
                         horizontal="center" if idx == 1 else col.align)
    ws.row_dimensions[row].height = HEIGHT_TOTAL


def _total_formula(letter: str, row_spec: RowSpec,
                   ref_to_row: dict[str, int]) -> str | None:
    """
    为「合计行」的某一列生成 SUM 公式，不是合计行则返回 None。

    两种写法：
        total_of   -> 连续区间，  =SUM(C8:C15)
        total_refs -> 跨段相加，  =SUM(C16,C29)
    传进来的是行标识 ref，这里负责翻译成真实的 Excel 行号。
    """
    if row_spec.kind != "total":
        return None

    if row_spec.total_of:
        r0 = ref_to_row.get(row_spec.total_of[0])
        r1 = ref_to_row.get(row_spec.total_of[1])
        if r0 and r1 and r1 >= r0:
            return f"=SUM({letter}{r0}:{letter}{r1})"
        return None

    if row_spec.total_refs:
        rows = [ref_to_row[r] for r in row_spec.total_refs if r in ref_to_row]
        if rows:
            return "=SUM(" + ",".join(f"{letter}{r}" for r in rows) + ")"
    return None


def _write_footer(ws, spec: ReportSpec, row: int, remark: bool) -> int:
    """写签字行和备注行，返回用到的最后一行行号。"""
    n_cols = len(spec.columns)
    last_letter = get_column_letter(n_cols)

    # 空一行，让签字区和表格正文分开
    ws.row_dimensions[row].height = 10
    row += 1

    if spec.footer_labels:
        ws.merge_cells(f"A{row}:{last_letter}{row}")
        cell = ws.cell(row=row, column=1)
        # 每个签字位后面留一段空白供手写签名
        cell.value = "        ".join(f"{label}　　　　　" for label in spec.footer_labels)
        cell.font = _font(size=10.5)
        cell.alignment = _align(horizontal="left")
        ws.row_dimensions[row].height = HEIGHT_FOOTER
        row += 1

    if remark and spec.remark_line:
        ws.merge_cells(f"A{row}:{last_letter}{row}")
        cell = ws.cell(row=row, column=1, value="备注：")
        cell.font = _font(size=10.5)
        cell.alignment = _align(horizontal="left", vertical="top")
        ws.row_dimensions[row].height = HEIGHT_FOOTER
        row += 1

    return row - 1


def _setup_print(ws, spec: ReportSpec, header_row: int, last_row: int) -> None:
    """设置打印参数：横向/纵向、缩放为一页宽、每页重复表头。"""
    ws.page_setup.orientation = "landscape" if spec.landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0        # 0 = 高度不限，按内容自动分页
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6,
                                  header=0.3, footer=0.3)
    # 跨页打印时每一页都重复表头行
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.print_area = f"A1:{get_column_letter(len(spec.columns))}{last_row}"


# ---------------------------------------------------------------------------
# 主入口
# ---------------------------------------------------------------------------

def build_workbook(spec: ReportSpec,
                   *,
                   base_bytes: bytes | None = None,
                   title: str = "",
                   org: str = "",
                   period: str = "",
                   detail_rows: int = 20,
                   with_total: bool = True,
                   with_remark: bool = True) -> Workbook:
    """
    根据报表定义生成一个排好版的空白 Excel 工作簿。

    参数：
        spec         : 报表定义（见 report_specs.py）
        base_bytes   : 用户上传的空白 Excel 的二进制内容；
                       传 None 则由引擎新建一个工作簿
        title        : 大标题文字，留空则用 spec.default_title
        org          : 编制单位名称
        period       : 期间，例如「2024 年度」或「2024 年 3 月」
        detail_rows  : 明细流水型报表预留的空白行数（固定行报表忽略此项）
        with_total   : 是否生成「合计」行
        with_remark  : 是否生成「备注：」行

    返回：
        openpyxl.Workbook 对象（尚未保存到磁盘）
    """
    wb, ws = _prepare_sheet(spec, base_bytes)

    header_row = _write_title_block(ws, spec, title, org, period)
    _write_header(ws, spec, header_row)

    first_data_row = header_row + 1
    row = first_data_row

    if spec.is_fixed_rows:
        # ---- 固定行报表：逐行写科目名，并为合计行解析真实行号 ----
        # 合计公式需要引用真实行号，而行号只有排版时才确定，
        # 所以先扫一遍建立「行标识 -> Excel 行号」的映射表
        ref_to_row: dict[str, int] = {
            rs.ref: first_data_row + i
            for i, rs in enumerate(spec.rows) if rs.ref
        }

        seq = 0    # 行次计数器：只给需要填数据的普通科目行编号
        for row_spec in spec.rows:
            is_section = row_spec.kind == "section"
            is_total = row_spec.kind == "total"
            bold = is_section or is_total
            if is_section:
                bg = COLOR_SECTION_BG
            elif is_total:
                bg = COLOR_TOTAL_BG
            else:
                bg = None

            if row_spec.kind == "item":
                seq += 1

            for idx, col in enumerate(spec.columns, start=1):
                cell = ws.cell(row=row, column=idx)

                if idx == 1:
                    # 第一列固定放科目名称，按层级缩进显示
                    cell.value = row_spec.label
                    _style_data_cell(cell, col, bg=bg, bold=bold,
                                     indent=row_spec.indent, horizontal="left")
                elif idx == spec.seq_column and spec.seq_column > 1:
                    # 行次列：分类行和合计行不占行次，留空
                    cell.value = seq if row_spec.kind == "item" else None
                    _style_data_cell(cell, col, bg=bg, bold=bold)
                else:
                    # 金额列：优先套用行定义里手写的公式，
                    # 其次才是合计行的 SUM；其他列留空待填
                    if col.summable:
                        letter = get_column_letter(idx)
                        if row_spec.formula:
                            try:
                                cell.value = row_spec.formula.format(
                                    col=letter, **ref_to_row)
                            except KeyError as exc:
                                # 公式模板里写了一个不存在的行标识，
                                # 直接报清楚的错，避免生成一份静默错误的表格
                                raise ValueError(
                                    f"报表「{spec.name}」第 {row} 行的公式引用了"
                                    f"不存在的行标识 {exc}，请检查 report_specs.py"
                                ) from None
                        else:
                            formula = _total_formula(letter, row_spec, ref_to_row)
                            if formula:
                                cell.value = formula
                    _style_data_cell(cell, col, bg=bg, bold=bold)

            ws.row_dimensions[row].height = (
                HEIGHT_TOTAL if is_total else
                HEIGHT_SECTION if is_section else HEIGHT_ITEM
            )
            row += 1

        last_data_row = row - 1
        # 科目里已经自带合计行时，不再额外追加一行「合计」
        if with_total and spec.has_total_row and not any(
                r.kind == "total" for r in spec.rows):
            _write_total_row(ws, spec, row, first_data_row, last_data_row)
            row += 1
    else:
        # ---- 明细流水型报表：铺 N 行带格式的空白行 ----
        detail_rows = max(1, int(detail_rows))
        for i in range(detail_rows):
            bg = COLOR_STRIPE_BG if (spec.stripe and i % 2 == 1) else None
            for idx, col in enumerate(spec.columns, start=1):
                cell = ws.cell(row=row, column=idx)
                # 序号列预先填好 1、2、3……，省得手工敲
                if idx == spec.seq_column:
                    cell.value = i + 1
                _style_data_cell(cell, col, bg=bg)
            ws.row_dimensions[row].height = HEIGHT_ITEM
            row += 1

        last_data_row = row - 1
        if with_total and spec.has_total_row:
            _write_total_row(ws, spec, row, first_data_row, last_data_row)
            row += 1
        # 下拉框只覆盖空白明细区，不覆盖合计行
        _add_dropdowns(ws, spec, first_data_row, last_data_row)

    # ---- 冻结窗格：滚动时表头始终可见 ----
    ws.freeze_panes = f"A{first_data_row}"

    # ---- 明细型报表加上筛选按钮，方便日后按条件查看 ----
    if not spec.is_fixed_rows and last_data_row >= first_data_row:
        last_letter = get_column_letter(len(spec.columns))
        ws.auto_filter.ref = f"A{header_row}:{last_letter}{last_data_row}"

    last_row = _write_footer(ws, spec, row, with_remark)
    _setup_print(ws, spec, header_row, last_row)

    # 文档属性，便于在文件详情里辨认
    wb.properties.title = spec.name
    wb.properties.creator = "报表格式生成器"
    wb.properties.description = f"由报表格式生成器创建：{spec.name}"

    return wb


def workbook_to_bytes(wb: Workbook) -> bytes:
    """把工作簿序列化成 .xlsx 的二进制内容（Streamlit 下载按钮要用）。"""
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
